"""Read plain text out of the file formats the app supports.

All readers return a single Unicode string. Structure-preserving round-trips
aren't a goal yet — anonymization runs on plain text, and the user re-exports
to a new file on save.
"""
from pathlib import Path

SUPPORTED_READ_SUFFIXES = (".txt", ".docx", ".xlsx", ".pdf")


class UnsupportedFormatError(ValueError):
    pass


def read_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return _read_txt(path)
    if suffix == ".docx":
        return _read_docx(path)
    if suffix == ".xlsx":
        return _read_xlsx(path)
    if suffix == ".pdf":
        return _read_pdf(path)
    raise UnsupportedFormatError(f"Формат не поддерживается: {suffix or '(без расширения)'}")


def _read_txt(path: Path) -> str:
    # Try UTF-8 first (most common), fall back to cp1251 for legacy Russian docs.
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "cp1251", "cp866"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _read_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    lines: list[str] = []
    for para in doc.paragraphs:
        lines.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            lines.append("\t".join(cells))
    return "\n".join(lines)


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True, read_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        parts.append(f"# {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(cells):
                parts.append("\t".join(cells))
        parts.append("")
    return "\n".join(parts).rstrip()


def _read_pdf(path: Path) -> str:
    import pdfplumber

    pages: list[str] = []
    with pdfplumber.open(str(path)) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            pages.append(f"--- Страница {i} ---\n{text}")
    return "\n\n".join(pages)
