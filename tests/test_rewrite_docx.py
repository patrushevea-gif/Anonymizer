from pathlib import Path

import pytest


def test_rewrite_docx_replaces_text_and_preserves_doc(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    from docx import Document

    from anonymizer.io import read_text, rewrite_docx

    src = tmp_path / "src.docx"
    dst = tmp_path / "dst.docx"

    doc = Document()
    doc.add_paragraph("Иван Иванов работает в ООО Рога и Копыта.")
    doc.add_paragraph("Его телефон +7 999 111 22 33, email ivan@example.com.")
    doc.save(str(src))

    rewrite_docx(
        src,
        dst,
        [
            ("Иван Иванов", "ФИО1"),
            ("ООО Рога и Копыта", "ОРГ1"),
            ("+7 999 111 22 33", "ТЕЛ1"),
            ("ivan@example.com", "EMAIL1"),
        ],
    )

    text = read_text(dst)
    assert "Иван Иванов" not in text
    assert "ФИО1" in text
    assert "ОРГ1" in text
    assert "ТЕЛ1" in text
    assert "EMAIL1" in text
