"""DOCX engine: every part processed, formatting kept, leaks verified."""
import io
import re
import zipfile
from pathlib import Path

import pytest

docx = pytest.importorskip("docx")


def _build(tmp_path: Path) -> Path:
    from docx.shared import RGBColor

    d = docx.Document()
    p = d.add_paragraph("Ответственный: ")
    r = p.add_run("Петров")
    r.bold = True
    r.font.color.rgb = RGBColor(0xC0, 0, 0)
    r2 = p.add_run(" Пётр Петрович")
    r2.italic = True
    p.add_run(", тел. +7 916 123-45-67.")

    hidden = d.add_paragraph("Видимо. ")
    h = hidden.add_run("Скрыто: ivan@example.com")
    h.font.hidden = True

    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "ИНН"
    t.cell(0, 1).text = "7707083893"

    sec = d.sections[0]
    sec.header.paragraphs[0].text = "Исполнитель: Петров Пётр Петрович"
    sec.footer.paragraphs[0].text = "Связь: ivan@example.com"

    commented = d.add_paragraph("Абзац с комментарием.")
    d.add_comment(commented.runs, text="Позвонить +7 916 123-45-67", author="Сидоров С.С.",
                  initials="СС")

    cp = d.core_properties
    cp.author = "Сидоров Семён Семёнович"
    cp.last_modified_by = "Сидоров Семён Семёнович"
    cp.title = "Договор с Петровым Петром Петровичем"
    cp.keywords = "ivan@example.com"
    path = tmp_path / "src.docx"
    d.save(str(path))
    return path


def _xml(data: bytes, name: str) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        return zf.read(name).decode("utf-8")


def _anonymize(path: Path):
    from anonymizer.core.mapping import MappingStore
    from anonymizer.core.pipeline import Pipeline
    from anonymizer.io.ooxml import DocxDocument

    doc = DocxDocument.open(path)
    store = MappingStore()
    result = Pipeline.default(use_ner=False).run(doc.text, store)
    data, report = doc.anonymize(result.spans, store)
    return doc, store, data, report


def test_text_includes_all_parts_with_labels(tmp_path: Path) -> None:
    from anonymizer.io.ooxml import DocxDocument

    text = DocxDocument.open(_build(tmp_path)).text
    assert "ИНН\t7707083893" in text  # table row kept in place, cells tab-separated
    for label in ("[Комментарии]", "[Верхний колонтитул]", "[Нижний колонтитул]",
                  "[Свойства документа]"):
        assert label in text
    assert "Позвонить +7 916 123-45-67" in text
    assert "Договор с Петровым Петром Петровичем" in text


def test_all_parts_masked_and_metadata_scrubbed(tmp_path: Path) -> None:
    pytest.importorskip("natasha")
    _, _, data, report = _anonymize(_build(tmp_path))
    everything = "".join(
        _xml(data, n) for n in zipfile.ZipFile(io.BytesIO(data)).namelist() if n.endswith(".xml")
    )
    for secret in ("Петров", "Пётр", "+7 916 123-45-67", "ivan@example.com", "7707083893",
                   "Сидоров", "Петровым"):
        assert secret not in everything, secret
    assert report.ok
    assert 'w:author="Автор"' in _xml(data, "word/comments.xml")


def test_formatting_and_hidden_text_preserved(tmp_path: Path) -> None:
    pytest.importorskip("natasha")
    _, _, data, _ = _anonymize(_build(tmp_path))
    body = _xml(data, "word/document.xml")
    # The alias inherits the bold red run where the name started.
    assert re.search(r"<w:b/>.*?<w:color w:val=\"C00000\"/>.*?<w:t[^>]*>ФИО1</w:t>", body, re.S)
    # Hidden run still hidden, now holding the alias.
    assert re.search(r"<w:vanish/>.*?<w:t[^>]*>Скрыто: EMAIL1</w:t>", body, re.S)


def test_restore_round_trip(tmp_path: Path) -> None:
    pytest.importorskip("natasha")
    from anonymizer.io.ooxml import DocxDocument

    doc, store, data, _ = _anonymize(_build(tmp_path))
    restored, count = DocxDocument(data).restore(store)
    assert count > 0
    assert DocxDocument(restored).text == doc.text


def test_unmodelled_location_is_scrubbed_by_leak_control(tmp_path: Path) -> None:
    """A value hiding in a part the engine doesn't model (custom XML) is still caught."""
    from anonymizer.core.mapping import MappingStore
    from anonymizer.core.pipeline import Pipeline
    from anonymizer.io.ooxml import DocxDocument

    src = _build(tmp_path)
    buf = io.BytesIO()
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(buf, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("customXml/item9.xml", "<root><owner>ivan@example.com</owner></root>")
    doc = DocxDocument(buf.getvalue())
    store = MappingStore()
    result = Pipeline.default(use_ner=False).run(doc.text, store)
    data, report = doc.anonymize(result.spans, store)
    assert "ivan@example.com" not in _xml(data, "customXml/item9.xml")
    assert ("customXml/item9.xml", "ivan@example.com") in report.leaks_fixed
    assert report.ok


def test_leak_is_reported_when_value_not_masked(tmp_path: Path) -> None:
    from anonymizer.core.entities import EntityKind
    from anonymizer.core.mapping import MappingStore
    from anonymizer.io.ooxml import DocxDocument

    doc = DocxDocument.open(_build(tmp_path))
    store = MappingStore()
    # Mask the phone in the body only; leak control must still clean the comment.
    start = doc.text.index("+7 916 123-45-67")
    from anonymizer.core.entities import EntitySpan

    span = EntitySpan(EntityKind.PHONE, start, start + 16, "+7 916 123-45-67", "+79161234567")
    data, report = doc.anonymize([span], store)
    assert "+7 916 123-45-67" not in _xml(data, "word/comments.xml")
    assert report.ok


def test_rewrite_docx_single_pass_no_chain_replacement(tmp_path: Path) -> None:
    from anonymizer.io import read_text, rewrite_docx

    d = docx.Document()
    d.add_paragraph("Сайт и IP: 192.0.2.47")
    src, dst = tmp_path / "a.docx", tmp_path / "b.docx"
    d.save(str(src))
    rewrite_docx(src, dst, [("192.0.2.47", "IP1"), ("IP", "ФОРМУЛА1")])
    text = read_text(dst)
    assert "IP1" in text and "ФОРМУЛА11" not in text


def test_batch_writes_docx_with_formatting(tmp_path: Path) -> None:
    pytest.importorskip("natasha")
    from anonymizer.core.pipeline import Pipeline
    from anonymizer.io.batch import collect_files, run_batch

    src_dir = tmp_path / "in"
    src_dir.mkdir()
    _build(src_dir)
    results = run_batch(Pipeline.default(use_ner=False), collect_files(src_dir), tmp_path / "out")
    assert results[0].error is None
    assert results[0].destination.suffix == ".docx"
    body = _xml(results[0].destination.read_bytes(), "word/header1.xml")
    assert "Петров" not in body


def test_strict_ooxml_is_refused_loudly(tmp_path: Path) -> None:
    from anonymizer.io.ooxml import DocxDocument

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr(
            "word/document.xml",
            '<w:document xmlns:w="http://purl.oclc.org/ooxml/wordprocessingml/main">'
            "<w:body><w:p><w:r><w:t>Иванов Иван Иванович</w:t></w:r></w:p></w:body></w:document>",
        )
    with pytest.raises(ValueError, match="Strict Open XML"):
        DocxDocument(buf.getvalue())
